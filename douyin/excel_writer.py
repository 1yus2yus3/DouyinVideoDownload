"""Excel 清单生成模块（openpyxl）。"""

from __future__ import annotations

import datetime
import re
from pathlib import Path

from .client import UserProfile, VideoItem


def _sanitize_filename(name: str, max_len: int = 40) -> str:
    """清理文件名中的非法字符并截断。"""
    name = re.sub(r'[\\/:*?"<>|\n\r\t]', "", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name[:max_len] or "untitled"


def build_excel(
    profile: UserProfile,
    videos: list[VideoItem],
    output_path: str | Path,
    extra_cols: bool = True,
) -> Path:
    """生成作品清单 Excel。

    extra_cols=True 时包含播放地址/封面等扩展列，False 时仅核心列。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "作品清单"

    # ---- 账号信息区 ----
    ws["A1"] = "账号信息"
    ws["A1"].font = Font(bold=True, size=13)
    info_rows = [
        ("账号昵称", profile.nickname),
        ("抖音号/ID", profile.sec_uid),
        ("签名", profile.signature),
        ("粉丝数", profile.follower_count),
        ("获赞数", profile.total_favorited),
        ("作品数", profile.aweme_count),
        ("抓取时间", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for i, (k, v) in enumerate(info_rows, start=2):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)

    # ---- 表头 ----
    headers = [
        "序号",
        "作品ID(aweme_id)",
        "标题",
        "发布时间",
        "时长(秒)",
        "mp4直链",
    ]
    if extra_cols:
        headers += ["作品页链接", "封面图", "全部播放地址"]

    header_row = len(info_rows) + 3  # 空一行后放表头
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # ---- 数据行 ----
    for idx, v in enumerate(videos, start=1):
        row = header_row + idx
        create_str = (
            datetime.datetime.fromtimestamp(v.create_time).strftime("%Y-%m-%d %H:%M")
            if v.create_time
            else ""
        )
        values = [
            idx,
            v.aweme_id,
            v.desc,
            create_str,
            v.duration_seconds,
            v.mp4_url,
        ]
        if extra_cols:
            values += [
                f"https://www.douyin.com/video/{v.aweme_id}",
                "; ".join(v.cover_urls[:1]),
                "\n".join(v.play_urls),
            ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=(col in (3, 9)))

    # ---- 列宽 ----
    widths = [6, 20, 50, 18, 10, 80, 30, 40, 90]
    for i, w in enumerate(widths[: len(headers)], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 冻结表头
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    wb.save(output_path)
    return output_path
